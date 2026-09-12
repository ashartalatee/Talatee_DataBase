"""
Helper baca CSV/XLSX jadi list-of-list baris mentah (termasuk header di
baris pertama) -- dipakai bersama oleh channel_guard.py dan
mixed_channel.py supaya logic parsing file tidak dobel dan tidak bisa
saling melenceng kalau salah satunya diubah nanti.
"""
import csv
import io

from openpyxl import load_workbook


def read_rows(raw_data: bytes, filename: str) -> list[list]:
    """Return [] kalau ekstensi tidak dikenali (biar pemanggil yang putuskan
    mau di-skip atau di-reject, tergantung konteks)."""
    ext = filename[filename.rfind(".") :].lower() if "." in filename else ""
    if ext == ".csv":
        return _read_csv_rows(raw_data)
    if ext == ".xlsx":
        return _read_excel_rows(raw_data)
    return []


# Nama kolom yang dikenali sebagai "channel/platform jualan" -- dipakai
# bareng oleh channel_guard.py (validasi 1 channel per file) dan
# mixed_channel.py (split 1 file jadi banyak channel).
CHANNEL_COLUMN_ALIASES = {"channel", "platform", "marketplace", "sumber"}


def normalize_channel_value(s) -> str:
    """Bandingkan longgar -- huruf/angka saja, tanpa spasi/tanda baca/besar-
    kecil huruf, supaya 'TikTok Shop', 'tiktokshop', 'TikTok-Shop' semua
    kehitung sama."""
    return "".join(ch for ch in str(s or "").lower() if ch.isalnum())


def _read_csv_rows(raw_data: bytes) -> list[list[str]]:
    text = raw_data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _read_excel_rows(raw_data: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(raw_data), read_only=True, data_only=True)
    sheet = wb.active
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    wb.close()
    return rows
