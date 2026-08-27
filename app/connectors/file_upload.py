"""
Connector pertama: menerima file yang SUDAH di-upload langsung (multipart),
bukan fetch dari sumber eksternal. Connector paling sederhana karena tidak
butuh koneksi eksternal apa pun — connect() jadi no-op (keputusan Phase 1 #2).

Mendukung minimal .csv dan .xlsx (Langkah 4 checklist).
"""
import csv
import io
from typing import Optional

from openpyxl import load_workbook

from app.connectors.base import Connector, FetchResult

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


class FileUploadConnector(Connector):
    def __init__(self, source_name: str, filename: str, file_bytes: bytes):
        self._source_name = source_name
        self._filename = filename
        self._file_bytes = file_bytes
        self._ext = self._get_extension(filename)
        if self._ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Ekstensi file '{self._ext}' tidak didukung. "
                f"Yang didukung: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

    @staticmethod
    def _get_extension(filename: str) -> str:
        idx = filename.rfind(".")
        return filename[idx:].lower() if idx != -1 else ""

    def connect(self) -> None:
        """No-op — file sudah tersedia di memory saat instance dibuat."""
        pass

    def fetch(self) -> FetchResult:
        if self._ext == ".csv":
            record_count, schema = self._parse_csv()
        else:  # .xlsx
            record_count, schema = self._parse_excel()

        return FetchResult(
            raw_data=self._file_bytes,
            filename=self._filename,
            content_type=self._content_type(),
            record_count=record_count,
            schema=schema,
        )

    def get_metadata(self) -> dict:
        connector_type = "file_upload_csv" if self._ext == ".csv" else "file_upload_excel"
        return {"source_name": self._source_name, "connector_type": connector_type}

    def _content_type(self) -> str:
        if self._ext == ".csv":
            return "text/csv"
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def _parse_csv(self) -> tuple[Optional[int], Optional[dict]]:
        """Parse minimal: hitung baris data (di luar header) + ambil nama kolom.
        Kalau file korup/tidak bisa dibaca, raise exception — akan ditangkap
        oleh IngestionEngine dan tercatat sebagai batch berstatus 'failed'."""
        try:
            text = self._file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(f"File CSV tidak bisa dibaca sebagai UTF-8: {exc}") from exc

        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            raise ValueError("File CSV kosong (tidak ada baris sama sekali).")

        header, data_rows = rows[0], rows[1:]
        schema = {"columns": header}
        return len(data_rows), schema

    def _parse_excel(self) -> tuple[Optional[int], Optional[dict]]:
        try:
            wb = load_workbook(io.BytesIO(self._file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"File Excel tidak bisa dibaca/korup: {exc}") from exc

        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration as exc:
            wb.close()
            raise ValueError("File Excel kosong (tidak ada baris sama sekali).") from exc

        data_row_count = sum(1 for _ in rows_iter)
        schema = {"columns": [str(c) if c is not None else "" for c in header]}
        wb.close()
        return data_row_count, schema
